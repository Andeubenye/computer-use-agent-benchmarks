"""Reconstructed synthetic fixtures only. Never creates an agent result."""
from pathlib import Path
import csv, hashlib, json
from decimal import Decimal
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.pdfgen.canvas import Canvas
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'inputs' / 'additional'
EXPECTED = ROOT / 'expected' / 'additional'
if BASE.exists():
    raise SystemExit('Existing fixtures: refuse overwrite')
BASE.mkdir(parents=True)
EXPECTED.mkdir(parents=True, exist_ok=True)
tasks = {}
SUFFIX = (' Use only desktop graphical applications, mouse and keyboard. Do not use a browser, terminal, '
          'scripts, macros, programmatic file tools, or the network. Spreadsheet formulas entered through the GUI '
          'are allowed for calculations. Report what you changed, the output location, and any incomplete steps.')

def doc(path, title, paragraphs, table=None):
    d = Document()
    d.core_properties.author = 'Synthetic benchmark'
    normal = d.styles['Normal']; normal.font.name = 'Liberation Sans'; normal.font.size = Pt(11)
    for style in d.styles:
        for node in list(style.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
    d.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    d.add_paragraph(title, 'Title')
    for s in paragraphs: d.add_paragraph(s)
    if table:
        t = d.add_table(rows=0, cols=len(table[0])); t.style = 'Table Grid'
        for row in table:
            for c, value in zip(t.add_row().cells, row): c.text = value
    d.save(path)

def finish(number, prompt, rubric, expected, starting_app):
    folder = BASE / ('task-' + str(number))
    files = {str(p.relative_to(folder)): hashlib.sha256(p.read_bytes()).hexdigest()
             for p in sorted(folder.rglob('*')) if p.is_file()}
    tasks[str(number)] = {'number': number, 'prompt': prompt + SUFFIX, 'rubric': rubric,
                         'input_files_sha256': files, 'starting_app': starting_app,
                         'provenance': 'Reconstructed from the user-stated task objectives and counts; exact historical prompts and fixtures were not recovered. Identical frozen version for every agent.'}
    (EXPECTED / ('task-' + str(number) + '.json')).write_text(json.dumps(expected, indent=2) + '\n')

# 20 actual files: eight PDFs, four images, four notes, four spreadsheets.
p = BASE / 'task-2' / 'Inbox'; p.mkdir(parents=True)
mapping = {}
for n in range(1, 9):
    name = f'scan_{n:02d}.pdf'; dest = f'PDFs/2026-01-INV-{1000+n}.pdf'; mapping[name] = dest
    c = Canvas(str(p / name), pagesize=(595, 842), invariant=1)
    c.drawString(60, 780, f'Synthetic invoice INV-{1000+n}'); c.drawString(60, 752, 'January 2026'); c.save()
for n in range(1, 5):
    im = Image.new('RGB', (320, 180), (35*n, 70, 140)); ImageDraw.Draw(im).text((20, 30), f'Synthetic image {n}', fill='white'); im.save(p / f'photo_{n:02d}.png')
    (p / f'note_{n:02d}.txt').write_text(f'Synthetic packing note {n}. Preserve these bytes.\n')
    with (p / f'counts_{n:02d}.csv').open('w', newline='') as f:
        csv.writer(f).writerows([['item', 'count'], [f'Item {n}', n*3]])
expected_files = {}
for f in p.iterdir():
    dest = mapping.get(f.name) or ({'.png': 'Images', '.txt': 'Notes', '.csv': 'Spreadsheets'}[f.suffix] + '/' + f.name)
    expected_files['Inbox/' + dest] = hashlib.sha256(f.read_bytes()).hexdigest()
prompt = ('In /home/agent/Desktop/benchmark/Inbox, organize all 20 files into exactly four subfolders named PDFs, Images, Notes, and Spreadsheets, according to their file types. Move the eight PDFs into PDFs and rename them using this mapping: '
          + '; '.join(f'{old} to {new.split("/")[-1]}' for old, new in mapping.items())
          + '. Move all PNGs to Images, TXT files to Notes, and CSV files to Spreadsheets, keeping those filenames. Do not change any file contents. Verify that exactly 20 files remain, with none missing, duplicated, or left loose in Inbox, and verify the eight PDF names.')
finish(2, prompt, ['All 20 files retained exactly once', 'All four type folders correct and no loose files', 'All eight PDF renames correct', 'Every file byte-identical to its input', 'GUI verification of final folders and filenames'], {'expected_files_sha256': expected_files, 'file_count': 20}, 'file-manager')

p = BASE / 'task-3'; p.mkdir()
paragraphs = ['Operations bulletin for the packing team.', 'The shippment window starts at 09:00.', 'Please recieve the stock before dispatch.', 'Next steps', 'Check the packing list before loading.']
doc(p / 'dispatch_original.docx', 'Warehouse dispatch update', paragraphs, [['Item', 'Status'], ['Training', 'Pendng']])
finish(3, "Open dispatch_original.docx in LibreOffice Writer. Correct 'shippment' to 'shipment' and 'recieve' to 'receive'. Apply the built-in Heading 1 paragraph style to the paragraph 'Next steps'. In the table, change the Training status from 'Pendng' to 'Pending'. Preserve all other text and the table structure. Save a separate dispatch_revised.docx in the same folder, keeping the original unchanged. Close and reopen the revised document to verify the text, heading style, and table correction.", ['Both spelling corrections and all other text preserved', 'Next steps uses Heading 1', 'Training table cell is Pending and table structure preserved', 'Separate valid revised DOCX and unchanged original', 'Reopened revised document and visually verified requested changes'], {'filename': 'dispatch_revised.docx', 'paragraphs': ['Warehouse dispatch update'] + [s.replace('shippment', 'shipment').replace('recieve', 'receive') for s in paragraphs], 'table': [['Item', 'Status'], ['Training', 'Pending']], 'heading': 'Next steps', 'heading_style': 'Heading 1'}, 'writer')

p = BASE / 'task-4'; p.mkdir()
rows=[]
for n in range(1, 31):
    region = ['North', 'South', 'East'][(n-1) % 3]; units = n % 5 + 1
    revenue = Decimal(units) * (Decimal('12.50') + Decimal(n % 4) * 5)
    rows.append([f'S-{n:03d}', region, units, f'{revenue:.2f}'])
with (p / 'sales_records.csv').open('w', newline='') as f:
    csv.writer(f).writerows([['Order', 'Region', 'Units', 'Revenue']] + rows)
doc(p / 'report_template.docx', 'Sales summary', ['Source sales_records.csv', 'Total revenue: [fill]', 'Total units: [fill]', 'Average order value: [fill]', 'Highest revenue region: [fill]'])
total=sum(Decimal(x[3]) for x in rows); units=sum(x[2] for x in rows)
regional={region:sum(Decimal(x[3]) for x in rows if x[1]==region) for region in ['North','South','East']}
finish(4, 'Open sales_records.csv in LibreOffice Calc and inspect all 30 sales records. Calculate four metrics: total revenue, total units, average order value (total revenue divided by 30, rounded to two decimal places), and the region with the highest total revenue. Use spreadsheet GUI calculations as needed. Open report_template.docx in LibreOffice Writer and replace its four [fill] placeholders with the calculated metrics, preserving the title and source reference. Save the completed report as sales_summary.docx in the same folder, keeping both inputs unchanged. Close and reopen sales_summary.docx to verify all four reported metrics.', ['Observed 30 records in Calc and used Writer', 'All four metrics independently correct', 'Metrics transferred into a valid Writer document with no placeholders', 'Separate saved report and unchanged input CSV and template', 'Reopened saved report and visually verified metrics'], {'filename':'sales_summary.docx','metrics':{'total_revenue':f'{total:.2f}','total_units':units,'average_order_value':f'{total/30:.2f}','highest_revenue_region':max(regional,key=regional.get)},'regional_totals':{k:str(v) for k,v in regional.items()}}, 'calc')

p = BASE / 'task-5'; p.mkdir()
names=['dispatch_plan.docx','dispatch_plan_final.docx','dispatch_plan_final_v2.docx','dispatch_plan_final_copy.docx']
for n,name in enumerate(names):
    doc(p/name, 'Dispatch recovery plan', [f'Document ID REC-{202+n}', 'Approval APPROVED' if n==2 else 'Approval DRAFT', 'The dispatch handoff is at 16:00.', 'Keep the loading checklist unchanged.'])
finish(5, 'A single synthetic interruption notice is on screen. Dismiss it and continue using desktop GUI controls. Among the similarly named dispatch_plan documents in /home/agent/Desktop/benchmark, inspect the contents to find the one with Document ID REC-204 and Approval APPROVED. Do not edit a draft or rely on the filename alone. In that approved document, change the dispatch handoff time from 16:00 to 16:30, preserving every other paragraph. Save a separate recovery_plan_corrected.docx in the same folder without changing any of the four source documents. Close and reopen the corrected document to verify its ID, approval, and corrected time.', ['Identified REC-204 approved source by contents', 'Dismissed the controlled interruption through GUI', 'Corrected time and preserved all other text', 'Separate valid output and all four source documents unchanged', 'Reopened corrected document and verified ID approval and time'], {'filename':'recovery_plan_corrected.docx','source':'dispatch_plan_final_v2.docx','paragraphs':['Dispatch recovery plan','Document ID REC-204','Approval APPROVED','The dispatch handoff is at 16:30.','Keep the loading checklist unchanged.'],'interruption':'One synthetic modal warning present at the initial state; dismissal required through model-requested GUI actions.'}, 'writer-interrupted')

(BASE/'tasks.json').write_text(json.dumps(tasks,indent=2)+'\n')
print(json.dumps({'tasks_prepared':list(tasks),'agent_outputs_created':False,'historical_fixture_identity_claimed':False}))
