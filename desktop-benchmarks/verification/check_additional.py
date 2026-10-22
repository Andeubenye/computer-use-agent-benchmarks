"""Disposable grader test cases, never agent outputs or claimed task results."""
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
import json, shutil
from verify_additional import ROOT, audit

spec=json.loads((ROOT/'inputs/additional/tasks.json').read_text())
passed=[]
with TemporaryDirectory() as tmp:
    base=Path(tmp)
    p=base/'organization';p.mkdir()
    expected=json.loads((ROOT/'expected/additional/task-2.json').read_text())['expected_files_sha256']
    sources={h:ROOT/'inputs/additional/task-2'/name for name,h in spec['2']['input_files_sha256'].items()}
    for name,digest in expected.items():
        dest=p/name;dest.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(sources[digest],dest)
    assert audit(2,p)['score']==4 and audit(2,p)['outcome']=='PARTIAL'
    passed.append('correct organization needs GUI verification')
    pdf=next((p/'Inbox/PDFs').glob('*.pdf'));shutil.copyfile(pdf,p/'duplicate.pdf')
    assert not list(audit(2,p)['requirements'].values())[0]
    passed.append('duplicate rejected');(p/'duplicate.pdf').unlink()
    pdf.write_bytes(b'altered')
    assert not list(audit(2,p)['requirements'].values())[3]
    passed.append('modified binary rejected')
    for task in [3,4,5]:
        folder=base/str(task);shutil.copytree(ROOT/f'inputs/additional/task-{task}',folder)
        truth=json.loads((ROOT/f'expected/additional/task-{task}.json').read_text())
        if task==3:
            d=Document(folder/'dispatch_original.docx')
            for paragraph,value in zip(d.paragraphs,truth['paragraphs']):paragraph.text=value
            next(p for p in d.paragraphs if p.text=='Next steps').style='Heading 1'
            d.tables[0].cell(1,1).text='Pending';d.save(folder/truth['filename'])
            assert audit(task,folder)['score']==4
            passed.append('document edits with missing reopen are partial')
            d.paragraphs[1].text='Unrequested change';d.save(folder/truth['filename'])
            assert not list(audit(task,folder)['requirements'].values())[0]
            passed.append('unrequested document edit rejected')
        elif task==4:
            d=Document(folder/'report_template.docx');m=truth['metrics']
            values=['Total revenue: $'+m['total_revenue'],'Total units: '+str(m['total_units']),'Average order value: $'+m['average_order_value'],'Highest revenue region: '+m['highest_revenue_region']]
            for paragraph,value in zip(d.paragraphs[2:],values):paragraph.text=value
            d.save(folder/truth['filename']);assert audit(task,folder)['score']==3
            passed.append('metrics checked independently without GUI claims')
            d.paragraphs[2].text='Total revenue: $17800';d.save(folder/truth['filename'])
            assert not list(audit(task,folder)['requirements'].values())[1]
            passed.append('incorrect arithmetic rejected')
        else:
            d=Document(folder/truth['source'])
            for paragraph,value in zip(d.paragraphs,truth['paragraphs']):paragraph.text=value
            d.save(folder/truth['filename']);assert audit(task,folder)['score']==3
            passed.append('recovery artifact needs interruption and reopen evidence')
            assert audit(task,folder,submitted=False)['score'] is None
            passed.append('unsubmitted recovery is not scored')
print(json.dumps({'offline_cases_passed':len(passed),'checks':passed,'provider_requests':0,'agent_outputs_created':False,'temporary_test_files_only':True}))
