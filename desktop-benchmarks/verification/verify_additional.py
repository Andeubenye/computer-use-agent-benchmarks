"""Read-only artifact audit. GUI evidence is graded separately after each run."""
from pathlib import Path
from collections import Counter
from docx import Document
import argparse, hashlib, json, re, csv
from decimal import Decimal

ROOT = Path(__file__).resolve().parents[1]
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def audit(task, folder, observations=None, submitted=True):
    folder=Path(folder); obs=observations or {}
    spec=json.loads((ROOT/'inputs/additional/tasks.json').read_text())[str(task)]
    expected=json.loads((ROOT/'expected/additional'/f'task-{task}.json').read_text())
    def gui(name):
        x=obs.get(name,{})
        return x.get('verified') is True and bool(x.get('evidence')) and all((ROOT/p).is_file() for p in x['evidence'])
    def originals():
        return all((folder/p).is_file() and sha(folder/p)==h for p,h in spec['input_files_sha256'].items())
    details={}; checks=[False]*5
    if not submitted:
        return {'outcome':'NOT TESTED','score':None,'task':task,'submitted':False,'checks':None}
    if task==2:
        files={str(p.relative_to(folder)):sha(p) for p in folder.rglob('*') if p.is_file() and not p.name.startswith('.')}
        truth=expected['expected_files_sha256']
        dirs={str(p.relative_to(folder)) for p in folder.rglob('*') if p.is_dir()}
        correct_dirs={'Inbox','Inbox/PDFs','Inbox/Images','Inbox/Notes','Inbox/Spreadsheets'}
        checks=[len(files)==20 and Counter(files.values())==Counter(truth.values()),set(files)==set(truth) and dirs==correct_dirs,all(p in files and files[p]==h for p,h in truth.items() if '/PDFs/' in p),Counter(files.values())==Counter(truth.values()),gui('final_folder_verification')]
        details={'actual_file_count':len(files),'actual_paths':sorted(files),'expected_paths':sorted(truth),'all_bytes_preserved':checks[3]}
        core=all(checks[:4])
    else:
        output=folder/expected['filename']; d=None
        if output.is_file():
            try:d=Document(output)
            except Exception as e:details['document_parse_error']=type(e).__name__
        paragraphs=[p.text for p in d.paragraphs] if d else []
        tables=[[[c.text for c in row.cells] for row in t.rows] for t in d.tables] if d else []
        details.update(output_exists=output.is_file(),valid_docx=d is not None,paragraphs=paragraphs,tables=tables,all_originals_preserved=originals())
        if task==3:
            text_ok=paragraphs==expected['paragraphs']
            heading=next((p for p in d.paragraphs if p.text==expected['heading']),None) if d else None
            style_ok=heading is not None and heading.style.name.lower().replace(' ','')=='heading1'
            checks=[text_ok,style_ok,tables==[expected['table']],d is not None and originals(),gui('reopened_and_verified')]
            details['heading_style']=heading.style.name if heading else None
            core=d is not None and text_ok and tables==[expected['table']]
        elif task==4:
            with (ROOT/'inputs/additional/task-4/sales_records.csv').open(newline='') as f:
                sales=list(csv.DictReader(f))
            total=sum(Decimal(row['Revenue']) for row in sales)
            regional={region:sum(Decimal(row['Revenue']) for row in sales if row['Region']==region) for region in {row['Region'] for row in sales}}
            recomputed={'total_revenue':f'{total:.2f}','total_units':sum(int(row['Units']) for row in sales),'average_order_value':f'{total/len(sales):.2f}','highest_revenue_region':max(regional,key=regional.get)}
            if recomputed!=expected['metrics']:raise ValueError('Expected sales metrics fail independent recomputation')
            details['independently_recomputed_metrics']=recomputed
            text='\n'.join(paragraphs+[' '.join(c for row in t for c in row) for t in tables])
            normalized=text.replace(',','')
            metrics={}
            patterns={'total_revenue':r'total\s+revenue\s*[:=]?\s*(?:USD\s*|\$\s*)?([0-9]+(?:\.[0-9]+)?)','total_units':r'total\s+units\s*[:=]?\s*([0-9]+)','average_order_value':r'average\s+order\s+value\s*[:=]?\s*(?:USD\s*|\$\s*)?([0-9]+(?:\.[0-9]+)?)','highest_revenue_region':r'highest\s+revenue\s+region\s*[:=]?\s*(North|South|East)'}
            for key,pattern in patterns.items():
                match=re.search(pattern,normalized,re.I)
                value=match.group(1) if match else None
                metrics[key]={'value':value,'correct':value is not None and (value.lower()==str(expected['metrics'][key]).lower() if key=='highest_revenue_region' else float(value)==float(expected['metrics'][key]))}
            correct=all(v['correct'] for v in metrics.values())
            checks=[gui('both_applications_and_records'),correct,d is not None and '[fill]' not in text and 'Sales summary' in text and 'sales_records.csv' in text,d is not None and originals(),gui('reopened_and_verified')]
            details['metrics']=metrics
            core=checks[1] and checks[2]
        else:
            identified='Document ID REC-204' in paragraphs and 'Approval APPROVED' in paragraphs
            text_ok=paragraphs==expected['paragraphs']
            checks=[identified,gui('interruption_dismissed'),text_ok,d is not None and originals(),gui('reopened_and_verified')]
            core=d is not None and identified and text_ok
    score=sum(checks)
    # Partial work remains partial, even when a mandatory core requirement is missing.
    outcome='PASS' if all(checks) else 'PARTIAL' if core or any(checks) else 'FAIL'
    return {'task':task,'outcome':outcome,'score':score,'requirements':dict(zip(spec['rubric'],checks)),'core_artifact_correct':core,'details':details,'gui_observations':obs,'response_used_as_proof':False,'verifier_edits_files':False}

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('task',type=int);p.add_argument('folder',type=Path);p.add_argument('--observations',type=Path);p.add_argument('--not-submitted',action='store_true');a=p.parse_args()
    obs=json.loads(a.observations.read_text()) if a.observations else None
    print(json.dumps(audit(a.task,a.folder,obs,not a.not_submitted),indent=2))
